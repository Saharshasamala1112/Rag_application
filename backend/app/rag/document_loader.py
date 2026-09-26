from __future__ import annotations

import csv
import io
import json
import re
import subprocess
import zipfile
from pathlib import Path

import pymupdf
import pymupdf4llm

from app.config.settings import settings


class DocumentLimitError(ValueError):
    """Raised when an upload exceeds a safe parser resource limit."""


def _check_file_size(file_path: Path, limit_mb: int, kind: str) -> None:
    if file_path.stat().st_size > limit_mb * 1024 * 1024:
        raise DocumentLimitError(
            f"{kind} exceeds the {limit_mb} MB processing limit."
        )


def normalize_text(text: str) -> str:
    """Normalize extracted text for consistent processing."""
    text = re.sub(r"[^\S\n]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = text.replace("\u2013", "-").replace("\u2014", " - ")
    text = text.replace("\u2018", "'").replace("\u2019", "'")
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    return text.strip()


SUPPORTED_EXTENSIONS = {
    ".pdf": "PDF Document",
    ".doc": "Word Document (Legacy)",
    ".docx": "Word Document",
    ".txt": "Plain Text",
    ".md": "Markdown",
    ".csv": "CSV Spreadsheet",
    ".json": "JSON Data",
    ".html": "HTML Document",
    ".htm": "HTML Document",
    ".xml": "XML Document",
    ".log": "Log File",
    ".py": "Python Source",
    ".js": "JavaScript Source",
    ".ts": "TypeScript Source",
    ".java": "Java Source",
    ".c": "C Source",
    ".cpp": "C++ Source",
    ".go": "Go Source",
    ".rs": "Rust Source",
    ".rb": "Ruby Source",
    ".php": "PHP Source",
    ".sh": "Shell Script",
    ".yaml": "YAML Config",
    ".yml": "YAML Config",
    ".toml": "TOML Config",
    ".ini": "INI Config",
    ".cfg": "Config File",
    ".conf": "Config File",
    ".sql": "SQL Query",
    ".r": "R Source",
    ".swift": "Swift Source",
    ".kt": "Kotlin Source",
    ".scala": "Scala Source",
    ".tex": "LaTeX Document",
    ".bib": "BibTeX Data",
    ".rtf": "Rich Text Format",
}


class DocumentLoader:
    """Load and extract markdown from various document formats."""

    @staticmethod
    def extract_markdown(file_path: Path) -> str:
        """Extract markdown text from a supported file format."""

        suffix = file_path.suffix.lower()

        if suffix == ".pdf":
            with pymupdf.open(file_path) as document:
                if document.page_count > settings.max_pdf_pages:
                    raise DocumentLimitError(
                        f"PDF exceeds the {settings.max_pdf_pages} page limit."
                    )
            raw = pymupdf4llm.to_markdown(str(file_path))
            if isinstance(raw, list):
                raw = str(raw)
            return normalize_text(raw)

        if suffix == ".docx":
            DocumentLoader._validate_docx_archive(file_path)
            return DocumentLoader._extract_docx(file_path)

        if suffix == ".doc":
            return DocumentLoader._extract_doc(file_path)

        if suffix == ".csv":
            _check_file_size(
                file_path,
                settings.max_tabular_file_mb,
                "CSV file",
            )
            return DocumentLoader._extract_csv(file_path)

        if suffix == ".json":
            _check_file_size(
                file_path,
                settings.max_tabular_file_mb,
                "JSON file",
            )
            return DocumentLoader._extract_json(file_path)

        if suffix in (".html", ".htm"):
            return DocumentLoader._extract_html(file_path)

        if suffix == ".xml":
            return DocumentLoader._extract_xml(file_path)

        if suffix == ".rtf":
            return DocumentLoader._extract_rtf(file_path)

        # Plain text and all code/config formats
        return DocumentLoader._extract_text(file_path)

    @staticmethod
    def _extract_text(file_path: Path) -> str:
        """Extract text from plain text files."""
        try:
            raw = file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            try:
                raw = file_path.read_text(encoding="latin-1")
            except Exception:
                raw = file_path.read_text(encoding="utf-8", errors="replace")
        return normalize_text(raw)

    @staticmethod
    def _extract_docx(file_path: Path) -> str:
        """Extract markdown from .docx files."""
        try:
            from docx import Document as DocxDocument

            doc = DocxDocument(str(file_path))
            lines: list[str] = []

            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue

                style_name = ""
                if para.style and para.style.name:
                    style_name = para.style.name

                if style_name.startswith("Heading 1"):
                    lines.append(f"\n# {text}\n")
                elif style_name.startswith("Heading 2"):
                    lines.append(f"\n## {text}\n")
                elif style_name.startswith("Heading 3"):
                    lines.append(f"\n### {text}\n")
                elif style_name.startswith("Heading 4"):
                    lines.append(f"\n#### {text}\n")
                elif style_name.startswith("Heading"):
                    lines.append(f"\n## {text}\n")
                elif style_name.startswith("List"):
                    lines.append(f"- {text}")
                else:
                    lines.append(text)

            # Extract tables
            for table in doc.tables:
                if not table.rows:
                    continue
                lines.append("")
                header = [cell.text.strip() for cell in table.rows[0].cells]
                lines.append("| " + " | ".join(header) + " |")
                lines.append("| " + " | ".join(["---"] * len(header)) + " |")
                for row in table.rows[1:]:
                    cells = [cell.text.strip() for cell in row.cells]
                    lines.append("| " + " | ".join(cells) + " |")
                lines.append("")

            return "\n".join(lines)

        except ImportError:
            # python-docx not installed, try antiword
            return DocumentLoader._extract_doc_antiword(file_path)
        except Exception:
            # Fallback: try reading as zip (docx is a zip)
            return DocumentLoader._extract_docx_fallback(file_path)

    @staticmethod
    def _validate_docx_archive(file_path: Path) -> None:
        max_archive_size = settings.max_document_archive_mb * 1024 * 1024
        max_document_xml_size = settings.max_docx_xml_mb * 1024 * 1024
        try:
            with zipfile.ZipFile(file_path) as archive:
                members = archive.infolist()
                total_size = sum(member.file_size for member in members)
                document_xml = archive.getinfo("word/document.xml")
        except (OSError, zipfile.BadZipFile, KeyError) as error:
            raise DocumentLimitError("DOCX archive is invalid.") from error

        if total_size > max_archive_size:
            raise DocumentLimitError(
                "DOCX archive exceeds the configured uncompressed size limit."
            )
        if document_xml.file_size > max_document_xml_size:
            raise DocumentLimitError(
                "DOCX document XML exceeds the configured size limit."
            )

    @staticmethod
    def _extract_doc(file_path: Path) -> str:
        """Extract text from legacy .doc files using antiword or catdoc."""
        return DocumentLoader._extract_doc_antiword(file_path)

    @staticmethod
    def _extract_doc_antiword(file_path: Path) -> str:
        """Extract .doc text using antiword or catdoc CLI."""
        # Try antiword first
        try:
            result = subprocess.run(
                ["antiword", str(file_path)],
                capture_output=True,
                text=True,
                timeout=settings.external_parser_timeout_seconds,
            )
            if result.returncode == 0 and result.stdout.strip():
                return result.stdout
        except (FileNotFoundError, subprocess.TimeoutExpired):
            pass

        # Try catdoc
        try:
            result = subprocess.run(
                ["catdoc", str(file_path)],
                capture_output=True,
                text=True,
                timeout=settings.external_parser_timeout_seconds,
            )
            if result.returncode == 0 and result.stdout.strip():
                return result.stdout
        except (FileNotFoundError, subprocess.TimeoutExpired):
            pass

        # Last resort: try reading as text
        try:
            content = file_path.read_bytes()
            # Try to find readable text in binary
            text = content.decode("latin-1")
            # Filter to printable ASCII + common chars
            cleaned = re.sub(r"[^\x20-\x7E\n\r\t]", " ", text)
            cleaned = re.sub(r"\s{3,}", "\n\n", cleaned)
            if len(cleaned.strip()) > 50:
                return cleaned.strip()
        except Exception:
            pass

        raise ValueError(
            f"Cannot extract text from '{file_path.name}'. "
            "Install 'antiword' or 'catdoc' for .doc support: "
            "sudo apt-get install antiword"
        )

    @staticmethod
    def _extract_docx_fallback(file_path: Path) -> str:
        """Fallback: extract raw XML text from docx (it's a zip)."""
        try:
            with zipfile.ZipFile(str(file_path), "r") as zf:
                with zf.open("word/document.xml") as f:
                    xml_content = f.read(
                        settings.max_docx_xml_mb * 1024 * 1024 + 1
                    ).decode("utf-8")
                    if len(xml_content.encode("utf-8")) > settings.max_docx_xml_mb * 1024 * 1024:
                        raise DocumentLimitError(
                            "DOCX document XML exceeds the configured size limit."
                        )
                    # Strip XML tags
                    text = re.sub(r"<[^>]+>", " ", xml_content)
                    text = re.sub(r"\s+", " ", text)
                    return text.strip()
        except DocumentLimitError:
            raise
        except Exception:
            return ""

    @staticmethod
    def _extract_csv(file_path: Path) -> str:
        """Extract markdown table from CSV."""
        try:
            text = file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = file_path.read_text(encoding="latin-1")

        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        if not rows:
            return ""

        lines = []
        header = rows[0]
        lines.append("| " + " | ".join(header) + " |")
        lines.append("| " + " | ".join(["---"] * len(header)) + " |")

        for row in rows[1:501]:
            padded = row + [""] * (len(header) - len(row))
            lines.append("| " + " | ".join(padded[: len(header)]) + " |")

        if len(rows) > 501:
            lines.append(f"\n*... {len(rows) - 501} more rows*")

        return "\n".join(lines)

    @staticmethod
    def _extract_json(file_path: Path) -> str:
        """Extract readable text from JSON."""
        try:
            text = file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = file_path.read_text(encoding="latin-1")

        try:
            data = json.loads(text)
            return json.dumps(data, indent=2, ensure_ascii=False)
        except json.JSONDecodeError:
            return text

    @staticmethod
    def _extract_html(file_path: Path) -> str:
        """Extract text from HTML, stripping tags."""
        try:
            text = file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = file_path.read_text(encoding="latin-1")

        text = re.sub(r"<script[^>]*>.*?</script>", "", text, flags=re.DOTALL | re.IGNORECASE)
        text = re.sub(r"<style[^>]*>.*?</style>", "", text, flags=re.DOTALL | re.IGNORECASE)
        text = re.sub(r"<(?:br|p|div|h[1-6]|li|tr)[^>]*/?>", "\n", text, flags=re.IGNORECASE)
        text = re.sub(r"<[^>]+>", "", text)
        text = text.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
        text = text.replace("&nbsp;", " ").replace("&quot;", '"')
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    @staticmethod
    def _extract_xml(file_path: Path) -> str:
        """Extract text from XML."""
        try:
            text = file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = file_path.read_text(encoding="latin-1")

        text = re.sub(r"<[^>]+>", " ", text)
        text = re.sub(r"\s+", " ", text)
        return text.strip()

    @staticmethod
    def _extract_rtf(file_path: Path) -> str:
        """Basic RTF text extraction."""
        try:
            text = file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            text = file_path.read_text(encoding="latin-1")

        text = re.sub(r"\\[a-zA-Z]+\d*\s?", "", text)
        text = text.replace("{", "").replace("}", "")
        text = re.sub(r"\\['\"]", "", text)
        return text.strip()
